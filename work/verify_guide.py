"""Validate delivered files and render local layout previews from workbook cells."""
from __future__ import annotations

import json
import math
import re
import zipfile
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

import openpyxl
from PIL import Image, ImageDraw, ImageFont

import guide_content as g
import guide_examples as examples
from build_guide import ROOT, OUT, STEM, font, wrapped_lines, MODE_NAMES


class FormulaEvaluator:
    """A small parser for the formula functions actually used in this workbook."""
    pattern=re.compile(r'"(?:[^"]|"")*"|\d+(?:\.\d+)?|[A-Z]+\d+|[A-Z_]+|>=|<=|<>|[=<>+*/(),-]')
    precedence={'=':1,'>=':1,'<=':1,'<>':1,'>':1,'<':1,'+':2,'-':2,'*':3,'/':3}

    def __init__(self,ws,overrides=None):
        self.ws=ws;self.overrides=overrides or {};self.cache={}

    def cell(self,name):
        if name in self.overrides:return self.overrides[name]
        if name in self.cache:return self.cache[name]
        value=self.ws[name].value
        result=self.formula(value) if isinstance(value,str) and value.startswith('=') else value
        self.cache[name]=result
        return result

    def formula(self,formula):
        text=formula[1:];tokens=self.pattern.findall(text)
        assert ''.join(tokens)==re.sub(r'\s+(?=(?:[^"]*"[^"]*")*[^"]*$)','',text),formula
        index=0
        def primary():
            nonlocal index
            token=tokens[index];index+=1
            if token.startswith('"'):return token[1:-1].replace('""','"')
            if re.fullmatch(r'\d+(\.\d+)?',token):return float(token) if '.' in token else int(token)
            if re.fullmatch(r'[A-Z]+\d+',token):return self.cell(token)
            if token=='(':
                value=expr(0);assert tokens[index]==')';index+=1;return value
            assert tokens[index]=='(',token
            index+=1;args=[]
            if tokens[index]!=')':
                while True:
                    args.append(expr(0))
                    if tokens[index]!=',':break
                    index+=1
            assert tokens[index]==')';index+=1
            if token=='IF':return args[1] if args[0] else args[2]
            if token=='OR':return any(args)
            if token=='MAX':return max(args)
            if token=='ROUNDUP':return math.ceil(args[0]*10**args[1])/10**args[1]
            raise AssertionError(token)
        def expr(minimum):
            nonlocal index
            lhs=primary()
            while index<len(tokens) and tokens[index] in self.precedence and self.precedence[tokens[index]]>=minimum:
                op=tokens[index];level=self.precedence[op];index+=1;rhs=expr(level+1)
                if op=='+':lhs+=rhs
                elif op=='-':lhs-=rhs
                elif op=='*':lhs*=rhs
                elif op=='/':lhs/=rhs
                elif op=='=':lhs=lhs==rhs
                elif op=='<>':lhs=lhs!=rhs
                elif op=='>=':lhs=lhs>=rhs
                elif op=='<=':lhs=lhs<=rhs
                elif op=='>':lhs=lhs>rhs
                elif op=='<':lhs=lhs<rhs
            return lhs
        answer=expr(0)
        assert index==len(tokens),(formula,tokens[index:])
        return answer


def validate_json(value):
    cs=value.get('containerDefinitions',[]) if isinstance(value,dict) else value
    if not isinstance(cs,list) or not cs or not isinstance(cs[0],dict) or 'name' not in cs[0]:return
    names={c['name'] for c in cs}
    assert len(names)==len(cs)==4
    edges={c['name']:[] for c in cs}
    by_name={c['name']:c for c in cs}
    for c in cs:
        h=c['healthCheck']
        assert 0<=h['startPeriod']<=300
        assert 5<=h['interval']<=300
        assert 2<=h['timeout']<=60
        assert 1<=h['retries']<=10
        assert h['command'][0] in ('CMD','CMD-SHELL')
        if 'startTimeout' in c:assert 2<=c['startTimeout']<=120
        assert 2<=c['stopTimeout']<=120
        assert c['essential']==(c['name']!='cwagent')
        for dep in c.get('dependsOn',[]):
            target=dep['containerName'];assert target in names
            assert dep['condition'] in ('START','HEALTHY','SUCCESS','COMPLETE')
            if dep['condition']=='HEALTHY':assert 'healthCheck' in by_name[target]
            if dep['condition'] in ('SUCCESS','COMPLETE'):assert not by_name[target]['essential']
            edges[c['name']].append(target)
    def visit(n,trail):
        assert n not in trail,edges
        for target in edges[n]:visit(target,trail+[n])
    for name in names:visit(name,[])


def actual_widths(ws):
    widths=[13]*ws.max_column
    for dim in ws.column_dimensions.values():
        for col in range(dim.min or 1,(dim.max or dim.min or 1)+1):
            if col<=len(widths):widths[col-1]=dim.width
    # OOXML width already includes the padding used by xlsxwriter.
    return [int(w*7) for w in widths]


def rgb(color,default):
    if color is None:return default
    if color.type=='rgb' and isinstance(color.rgb,str):return '#'+color.rgb[-6:]
    if color.type=='theme':return '#FFFFFF' if color.theme==0 else default
    return default


def render(ws,path,start=1,end=None,max_width=1600):
    end=min(ws.max_row,end or ws.max_row)
    widths=actual_widths(ws)
    xs=[0]
    for w in widths:xs.append(xs[-1]+w)
    heights={r:math.ceil((ws.row_dimensions[r].height or 15)*96/72) for r in range(start,end+1)}
    ys={start:0}
    for r in range(start,end+1):ys[r+1]=ys[r]+heights[r]
    im=Image.new('RGB',(xs[-1]+2,ys[end+1]+2),'white');draw=ImageDraw.Draw(im)
    merged={}
    skip=set()
    for ran in ws.merged_cells.ranges:
        merged[(ran.min_row,ran.min_col)]=(ran.max_row,ran.max_col)
        for row in range(ran.min_row,ran.max_row+1):
            for col in range(ran.min_col,ran.max_col+1):
                if (row,col)!=(ran.min_row,ran.min_col):skip.add((row,col))
    for row in range(start,end+1):
        for col in range(1,ws.max_column+1):
            if (row,col) in skip:continue
            cell=ws.cell(row,col)
            endrow,endcol=merged.get((row,col),(row,col))
            if endrow>end:continue
            x,y=xs[col-1],ys[row];x2,y2=xs[endcol],ys[endrow+1]
            fill=rgb(cell.fill.fgColor,'#FFFFFF') if cell.fill.patternType=='solid' else '#FFFFFF'
            draw.rectangle((x,y,x2,y2),fill=fill)
            if any(side.style for side in [cell.border.left,cell.border.right,cell.border.top,cell.border.bottom]):draw.rectangle((x,y,x2,y2),outline='#C5C5C5',width=1)
            if cell.value is None:continue
            text=str(cell.value);size=cell.font.sz or 11;f=font(size,bool(cell.font.b))
            color=rgb(cell.font.color,'#222222')
            lines=wrapped_lines(text,x2-x,size,bool(cell.font.b)) if cell.alignment.wrap_text else text.split('\n')
            lineh=round(size*1.55*96/72)
            text_h=len(lines)*lineh
            yy=y+6 if cell.alignment.vertical!='center' else y+max(3,(y2-y-text_h)/2)
            for line in lines:
                xx=x+7
                if cell.alignment.horizontal=='center':xx=x+max(3,(x2-x-f.getlength(line))/2)
                draw.text((xx,yy),line,font=f,fill=color)
                yy+=lineh
    if im.width>max_width:
        ratio=max_width/im.width;im=im.resize((max_width,round(im.height*ratio)),Image.Resampling.LANCZOS)
    im.save(path)


def check_sources():
    def fetch(src):
        sid,title,url,claim=src
        try:
            with urlopen(Request(url,headers={'User-Agent':'Mozilla/5.0'},method='HEAD'),timeout=20) as response:
                return {'id':sid,'status':response.status,'url':response.url}
        except (HTTPError,URLError,TimeoutError) as exc:return {'id':sid,'error':str(exc)}
    with ThreadPoolExecutor(max_workers=4) as pool:
        return sorted(pool.map(fetch,g.SOURCES),key=lambda item:item['id'])


def main():
    xlsx=OUT/(STEM+'.xlsx');md=OUT/(STEM+'.md')
    w=openpyxl.load_workbook(xlsx,data_only=False)
    cached=openpyxl.load_workbook(xlsx,data_only=True)
    manifest=json.loads((ROOT/'work/build_manifest.json').read_text(encoding='utf-8'))
    assert zipfile.ZipFile(xlsx).testzip() is None
    assert len(w.sheetnames)==21
    assert '\ufffd' not in md.read_text(encoding='utf-8')
    assert not md.read_bytes().startswith(b'\xef\xbb\xbf')
    colors=set();cell_count=0
    formula_count=0
    for ws in w:
        assert ws.sheet_view.showGridLines is False
        assert ws.freeze_panes is not None
        assert ws.print_area is not None
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        for row in ws:
            for c in row:
                if c.value is None:continue
                cell_count+=1
                assert c.font.name=='Meiryo UI',(ws.title,c.coordinate,c.font.name)
                assert '\ufffd' not in str(c.value)
                assert len(str(c.value))<=32767
                for color in [c.font.color,c.fill.fgColor if c.fill.patternType=='solid' else None]:
                    if color is not None and color.type=='rgb' and isinstance(color.rgb,str):
                        v=color.rgb[-6:];colors.add(v);assert v[:2]==v[2:4]==v[4:],v
                if c.data_type=='f':
                    formula_count+=1
                    assert cached[ws.title][c.coordinate].value is not None
    calc=manifest['calculator'];cs=calc['cells'];ws=w[calc['sheet']]
    ev=FormulaEvaluator(ws)
    for result in calc['results']:
        assert ev.cell(result['cell'])==result['cached'],result
        assert ev.cell(result['cell'])==cached[ws.title][result['cell']].value
    cases=[]
    for z in g.PATTERNS:
        t=g.timing(z)
        values={cs['mode']:MODE_NAMES[z['mode']],cs['a']:z['a'],cs['b']:z['b'],cs['f']:z['f'],cs['da']:z['da'],cs['db']:z['db'],cs['df']:z['df'],cs['m']:z['margin'],cs['g']:z['grace'],cs['ts']:z['task'] or 0}
        e=FormulaEvaluator(ws,values)
        for key in ['ha','hb','hf','bs','fs','jvm','alb','end']:assert e.cell(cs[key])==t[key],(z['id'],key)
        expected='依存構成の見直し' if z['id']=='P3' else '仮定上は猶予内'
        assert e.cell(cs['overall'])==expected
        if z['task'] is not None:assert e.cell(cs['need'])==t['need']
        cases.append({'pattern':z['id'],'result':expected,'need':None if z['id']=='P3' else t['need']})
    edgecases=[
        ({cs['mode']:'直列',cs['b']:110,cs['db']:10},'依存構成の見直し'),
        ({cs['mode']:'直列',cs['b']:109,cs['db']:10,cs['g']:600},'仮定上は猶予内'),
        ({cs['g']:0},'猶予不足'),
        ({cs['a']:120,cs['da']:0},'依存構成の見直し'),
    ]
    for vals,expected in edgecases:assert FormulaEvaluator(ws,vals).cell(cs['overall'])==expected
    assert len(ws.data_validations.dataValidation)==16
    json_paths=list((ROOT/'work/examples').glob('*.json'))
    for path in json_paths:validate_json(json.loads(path.read_text(encoding='utf-8')))
    fenced=re.findall(r'```json\n(.*?)\n```',md.read_text(encoding='utf-8'),flags=re.S)
    for snippet in fenced:validate_json(json.loads(snippet))
    source_ids=set(SOURCE_ID[0] for SOURCE_ID in g.SOURCES)
    for s in g.SECTIONS:
        for b in s['blocks']:assert set(b.get('refs','').split())<=source_ids
    preview=ROOT/'work/previews';preview.mkdir(exist_ok=True)
    render(cached['00_結論と読み方'],preview/'00_overview.png',1,25)
    render(cached['04_パターン比較'],preview/'04_patterns.png',1,18)
    render(cached['05_コンテナ設定'],preview/'05_settings.png',1,24)
    render(cached['07b_時系列図'],preview/'07_timeline.png')
    render(cached['08_時間計算'],preview/'08_calculator.png',34)
    render(cached['12_JSON設定例'],preview/'12_json.png',5,25)
    report={
        'xlsx_zip_integrity':'passed','sheet_count':len(w.sheetnames),'filled_cells':cell_count,
        'all_text_font':'Meiryo UI','monochrome_colors':sorted(colors),'formula_count':formula_count,
        'formula_checks':'14 cached results + 8 scenarios + 4 boundary cases passed',
        'scenario_results':cases,'input_validation_count':16,'json_template_count':len(json_paths),
        'markdown_json_blocks':len(fenced),'source_count':len(g.SOURCES),
        'markdown_characters':len(md.read_text(encoding='utf-8')),'preview_count':6,
        'native_excel_render':'not available: COM server startup failed; local cell-layout previews used',
        'aws_execution':'not performed; no AWS changes',
        'terraform_validate_plan':'not performed: illustrative fragments lack deployment-specific provider/variables/resources',
        'probe_binary_build':'not performed: Go toolchain and target image not supplied',
    }
    (ROOT/'work/validation_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
